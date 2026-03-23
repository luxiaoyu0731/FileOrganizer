from pathlib import Path

MAX_CHARS = 2000


def extract_docx(path: Path) -> dict:
    try:
        import docx
        doc = docx.Document(str(path))
        props = doc.core_properties
        metadata = {
            "title": props.title or "",
            "author": props.author or "",
        }
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return {"summary_text": text[:MAX_CHARS], "metadata": metadata}
    except Exception as e:
        return {"summary_text": "", "metadata": {"extraction_error": str(e)}}


def extract_xlsx(path: Path) -> dict:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines = [f"工作表：{', '.join(wb.sheetnames)}"]
        for sheet_name in wb.sheetnames[:3]:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=10, values_only=True))
            for row in rows:
                row_str = " | ".join(str(c) for c in row if c is not None)
                if row_str.strip():
                    lines.append(row_str)
        wb.close()
        return {"summary_text": "\n".join(lines)[:MAX_CHARS], "metadata": {"sheets": wb.sheetnames}}
    except Exception as e:
        return {"summary_text": "", "metadata": {"extraction_error": str(e)}}


def extract_pptx(path: Path) -> dict:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        titles = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and shape.shape_type == 13:
                    continue
                if hasattr(shape, "text") and shape.text.strip():
                    titles.append(shape.text.strip())
                    break
        metadata = {"slide_count": len(prs.slides)}
        return {"summary_text": "\n".join(titles)[:MAX_CHARS], "metadata": metadata}
    except Exception as e:
        return {"summary_text": "", "metadata": {"extraction_error": str(e)}}
